import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.opc.constants import RELATIONSHIP_TYPE
import os

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text=None, color_hex="0D6EFD"):
    if text is None:
        text = url
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:color w:val="{color_hex}"/><w:u w:val="single"/><w:sz w:val="19"/></w:rPr><w:t>{text}</w:t></w:r>')
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_document(output_path):
    doc = docx.Document()

    # Set page margins (narrow/standard: 2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header / Brand
    p_brand = doc.add_paragraph()
    p_brand.paragraph_format.space_before = Pt(0)
    p_brand.paragraph_format.space_after = Pt(2)
    run_brand = p_brand.add_run("CSA BOOKS 4 KIDS")
    run_brand.font.name = "Calibri"
    run_brand.font.size = Pt(10)
    run_brand.font.bold = True
    run_brand.font.color.rgb = RGBColor(240, 90, 40) # Brand orange #F05A28

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("Tabella Link Ufficiali Bonus & QR Code")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(21, 29, 42) # Slate black #151D2A

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Riepilogo delle pagine e dei contenuti bonus per le raccolte cartacee (Volume 1 e Volume 2), con collegamenti diretti in chiaro.")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    tables_data = [
        {
            "title": "📚 Raccolta Volume 1",
            "asin": "B0HN4QMLV1",
            "amazon_url": "https://www.amazon.it/dp/B0HN4QMLV1",
            "accent_color": "F05A28",
            "rows": [
                (
                    "Raccolta Vol. 1 — Disegni da Colorare",
                    "PDF stampabile (invio automatico via Brevo API)",
                    "https://csabookskids.com/csa-vol1-bonus-8f2e91b7d4a6"
                ),
                (
                    "Storia 1: Benny l'escavatore e la collina che cambiava forma",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol1-ep1-ita-7b3f91a8d2c4"
                ),
                (
                    "Storia 2: Rudy la ruspa e la spiaggia che scompariva",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol1-ep2-ita-4d9e16a2c8f1"
                ),
                (
                    "Storia 3: Leo la gru e le gemme della montagna",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol1-ep3-ita-9a5c83e1f7d2"
                ),
                (
                    "Storia 4: Bruno il camion e lo stagno delle papere",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol1-ep4-ita-6e2b94f0a3c5"
                ),
            ]
        },
        {
            "title": "🏗️ Raccolta Volume 2",
            "asin": "B0HN9G1V81",
            "amazon_url": "https://www.amazon.it/dp/B0HN9G1V81",
            "accent_color": "1E3A8A", # Dark Navy
            "rows": [
                (
                    "Raccolta Vol. 2 — Disegni da Colorare",
                    "PDF stampabile (invio automatico via Brevo API)",
                    "https://csabookskids.com/csa-vol2-bonus-46c694dc3ec4"
                ),
                (
                    "Storia 1: Benny l'escavatore e le monete dell'antica Roma",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol2-ep1-ita-1c683f8acb67"
                ),
                (
                    "Storia 2: Nina la betoniera e il ponte dei coniglietti",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol2-ep2-ita-ed3dabb448ca"
                ),
                (
                    "Storia 3: Rino il rullo e i ladri del cantiere",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol2-ep3-ita-2779f3736b81"
                ),
                (
                    "Storia 4: Leo la gru e il nido della cinciallegra",
                    "Video Read-Aloud + Traccia Audio MP3/WAV",
                    "https://csabookskids.com/csa-vol2-ep4-ita-605c1bd7e4ce"
                ),
            ]
        }
    ]

    for section_idx, t_data in enumerate(tables_data):
        if section_idx > 0:
            doc.add_paragraph().paragraph_format.space_before = Pt(16)

        # Section Heading
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(10)
        p_sec.paragraph_format.space_after = Pt(4)
        run_sec = p_sec.add_run(t_data["title"])
        run_sec.font.name = "Calibri"
        run_sec.font.size = Pt(14)
        run_sec.font.bold = True
        run_sec.font.color.rgb = RGBColor(21, 29, 42)

        p_meta = doc.add_paragraph()
        p_meta.paragraph_format.space_before = Pt(0)
        p_meta.paragraph_format.space_after = Pt(8)
        run_meta = p_meta.add_run(f"ASIN: {t_data['asin']}  |  Amazon: ")
        run_meta.font.name = "Calibri"
        run_meta.font.size = Pt(10)
        run_meta.font.color.rgb = RGBColor(71, 85, 105)
        add_hyperlink(p_meta, t_data["amazon_url"], t_data["amazon_url"], color_hex="F05A28")

        # Table
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        col_widths = [Inches(2.5), Inches(1.8), Inches(2.6)]

        # Header Row
        hdr_cells = table.rows[0].cells
        hdr_titles = ["Elemento / Contenuto", "Tipo di Bonus", "URL Diretto (in chiaro)"]
        for i, title in enumerate(hdr_titles):
            hdr_cells[i].width = col_widths[i]
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(title)
            run.font.name = "Calibri"
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            set_cell_shading(hdr_cells[i], t_data["accent_color"])
            set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
            hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Data Rows
        for r_idx, row_data in enumerate(t_data["rows"]):
            row_cells = table.add_row().cells
            bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"

            for i in range(3):
                row_cells[i].width = col_widths[i]
                set_cell_shading(row_cells[i], bg_color)
                set_cell_margins(row_cells[i], top=120, bottom=120, left=160, right=160)
                row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            # Col 1: Title
            p0 = row_cells[0].paragraphs[0]
            p0.paragraph_format.space_before = Pt(1)
            p0.paragraph_format.space_after = Pt(1)
            run0 = p0.add_run(row_data[0])
            run0.font.name = "Calibri"
            run0.font.size = Pt(9.5)
            run0.font.bold = True
            run0.font.color.rgb = RGBColor(30, 41, 59)

            # Col 2: Bonus type
            p1 = row_cells[1].paragraphs[0]
            p1.paragraph_format.space_before = Pt(1)
            p1.paragraph_format.space_after = Pt(1)
            run1 = p1.add_run(row_data[1])
            run1.font.name = "Calibri"
            run1.font.size = Pt(9.0)
            run1.font.color.rgb = RGBColor(71, 85, 105)

            # Col 3: Link in chiaro
            p2 = row_cells[2].paragraphs[0]
            p2.paragraph_format.space_before = Pt(1)
            p2.paragraph_format.space_after = Pt(1)
            add_hyperlink(p2, row_data[2], row_data[2], color_hex="0284C7")

        # Table borders: subtle border
        tblPr = table._tbl.tblPr
        tblBorders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
            f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(tblBorders)

    # Footer note
    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(10)
    p_note.paragraph_format.space_after = Pt(4)
    run_note = p_note.add_run("⚠️ Nota di integrità: ")
    run_note.font.name = "Calibri"
    run_note.font.size = Pt(9)
    run_note.font.bold = True
    run_note.font.color.rgb = RGBColor(220, 38, 38)

    run_note_text = p_note.add_run(
        "Tutti i percorsi e gli slug sopra elencati sono stampati in formato QR Code fisico "
        "all'interno dei libri distribuiti su Amazon. Non devono essere mai rinominati o spostati."
    )
    run_note_text.font.name = "Calibri"
    run_note_text.font.size = Pt(9)
    run_note_text.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(output_path)
    print(f"File salvato con successo in: {output_path}")

if __name__ == "__main__":
    out_file = os.path.join(os.getcwd(), "Tabelle_Bonus_CSA_Volume_1_e_2.docx")
    create_document(out_file)
